from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables" / "sales-kit" / "Nexatech School Pilot Proposal and Budget.docx"
SHOTS = ROOT / "deliverables" / "sales-kit" / "screenshots"

NAVY = "14385F"
BLUE = "0B6FA4"
TEAL = "1B8A72"
PALE = "EEF4F8"
LIGHT = "F5F7F8"
BORDER = "D6DEE5"
TEXT = "17232E"
MUTED = "5B6772"


def font(run, size=10.3, bold=False, color=TEXT, name="Aptos"):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), name)
    rpr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def space(p, before=0, after=6, line=1.1):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line


def body(doc, text, lead=None, after=7):
    p = doc.add_paragraph()
    space(p, after=after, line=1.12)
    if lead and text.startswith(lead):
        font(p.add_run(lead), bold=True)
        font(p.add_run(text[len(lead):]))
    else:
        font(p.add_run(text))
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    font(p.add_run(text))
    space(p, after=3, line=1.08)
    return p


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margins(cell, value=105):
    tcpr = cell._tc.get_or_add_tcPr()
    tcmar = tcpr.first_child_found_in("w:tcMar")
    if tcmar is None:
        tcmar = OxmlElement("w:tcMar")
        tcpr.append(tcmar)
    for side in ("top", "start", "bottom", "end"):
        item = tcmar.find(qn(f"w:{side}"))
        if item is None:
            item = OxmlElement(f"w:{side}")
            tcmar.append(item)
        item.set(qn("w:w"), str(value))
        item.set(qn("w:type"), "dxa")


def table_borders(table):
    pr = table._tbl.tblPr
    borders = pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), BORDER)


def repeat_header(row):
    pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    pr.append(node)


def style_cell(cell, bold=False, color=TEXT, size=9.1):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    cell_margins(cell)
    for p in cell.paragraphs:
        space(p, after=0, line=1.04)
        for run in p.runs:
            font(run, size=size, bold=bold, color=color)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table_borders(table)
    repeat_header(table.rows[0])
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = Inches(widths[i])
        cell.text = header
        shade(cell, NAVY)
        style_cell(cell, bold=True, color="FFFFFF", size=8.8)
    for ri, values in enumerate(rows):
        row = table.add_row()
        for ci, value in enumerate(values):
            cell = row.cells[ci]
            cell.width = Inches(widths[ci])
            cell.text = str(value)
            if ri % 2:
                shade(cell, LIGHT)
            style_cell(cell)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_image(doc, file_name, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture = p.add_run().add_picture(str(SHOTS / file_name), width=Inches(6.9))
    picture._inline.docPr.set("descr", caption)
    picture._inline.docPr.set("title", caption)
    space(p, after=3)
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(c.add_run(caption), size=8.7, color=MUTED)
    space(c, after=8)


doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.65)
sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.72)
sec.right_margin = Inches(0.72)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(10.3)
normal.font.color.rgb = RGBColor.from_string(TEXT)
normal.paragraph_format.line_spacing = 1.1
normal.paragraph_format.space_after = Pt(6)

title = styles["Title"]
title.font.name = "Aptos Display"
title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
title.font.size = Pt(25)
title.font.bold = True
title.font.color.rgb = RGBColor(0, 0, 0)
title.paragraph_format.space_after = Pt(8)

for name, size in (("Heading 1", 16), ("Heading 2", 12.5)):
    st = styles[name]
    st.font.name = "Aptos Display"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 9)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.keep_with_next = True

header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
font(header.add_run("NEXATECH SCHOOL OPERATIONS PLATFORM"), size=8, bold=True, color=MUTED)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(footer.add_run("School Pilot Proposal and Budget  |  21 October 2026"), size=8, color=MUTED)

doc.add_paragraph("Nexatech School Pilot Proposal and Budget", style="Title")
p = doc.add_paragraph()
space(p, after=12)
font(p.add_run("A 30-day evaluation of student records, marks, results, and attendance"), size=12, bold=True, color=BLUE)

body(doc, "Nexatech is a school operations platform that brings routine administrative and academic work into one controlled system. Staff capture information once, authorized users retrieve it when needed, and the school keeps a clearer operational record.")
body(doc, "This proposal recommends a focused 30-day pilot. The school can test the system using an agreed data set, named staff, measurable workflows, and a defined budget before deciding on a wider rollout.")

doc.add_heading("Recommended Offer", level=1)
add_table(doc,
    ["Item", "Scope", "Recommended price"],
    [
        ("30-day pilot", "Configuration, one clean import, one administrator, up to 10 teachers, training, 30-day pilot hosting, database, backups, and support", "US$299 once"),
        ("Academic plan", "Up to 500 active learners, records, Fast Marks, assessments, results, report cards, attendance, and exports", "US$179 per month"),
        ("Operations plan", "Academic plan plus finance, timetable, communications, behaviour, parent portal, and priority support", "US$299 per month"),
        ("Public API access", "One documented integration, scoped credentials, logs, standard rate limit, and key rotation support", "US$49 per month"),
        ("Managed domain and hosting", "Railway deployment, custom domain, HTTPS, standard service monitoring, and routine platform maintenance", "US$50 per month"),
    ],
    [1.35, 4.45, 1.35]
)
body(doc, "Pilot fee credit: if the school signs the Academic or Operations plan within 14 days of the pilot review, apply the US$299 pilot fee toward the standard US$350 onboarding charge.", "Pilot fee credit:")
body(doc, "Separate charges: SMS, email, payment-provider charges, taxes, complex historical data cleaning, custom integrations, and dedicated database resources are quoted separately.", "Separate charges:")

doc.add_heading("What the School Will See", level=1)
for item in [
    "An operational dashboard showing student and billing information.",
    "A searchable student register with role-controlled actions.",
    "Fast Marks for class, subject, term, assessment, and roster-based entry.",
    "Class results and rankings that staff can retrieve by term and year.",
    "Attendance registers with status and remarks.",
    "A route into finance, timetables, communications, behaviour, examinations, report cards, and parent access.",
]:
    bullet(doc, item)

add_image(doc, "01-dashboard.png", "Operational dashboard using synthetic Greenwood Academy demonstration data")

doc.add_heading("Pilot Demonstration Agenda", level=1)
add_table(doc,
    ["Time", "Demonstration", "Decision point"],
    [
        ("5 minutes", "Current school process and agreed pilot problem", "Confirm the workflow that matters most"),
        ("5 minutes", "Role-based login, navigation, and student register", "Confirm users and access boundaries"),
        ("10 minutes", "Fast Marks entry and safe retrieval", "Confirm required assessment process"),
        ("5 minutes", "Class results and attendance", "Confirm reporting and register expectations"),
        ("5 minutes", "Pilot scope, budget, data preparation, and next steps", "Nominate school lead and pilot dates"),
    ],
    [1.0, 3.65, 2.5]
)

add_image(doc, "03-fast-marks.png", "Fast Marks demonstration with synthetic class and assessment records")

doc.add_heading("Pilot Plan", level=1)
add_table(doc,
    ["Period", "Nexatech work", "School participation"],
    [
        ("Week 1", "Configure the workspace and import one clean data file", "Approve scope, data fields, named users, and school lead"),
        ("Week 2", "Train the administrator and teachers", "Attend training and confirm the operating process"),
        ("Week 3", "Support student records, marks, results, and attendance", "Use the agreed workflows and report issues"),
        ("Week 4", "Review outcomes and recommend the next rollout", "Assess value, risks, and the continuation decision"),
    ],
    [1.0, 3.4, 2.75]
)

doc.add_heading("Pilot Success Measures", level=1)
for item in [
    "A teacher can complete a 40-learner assessment in under five minutes.",
    "Common student and results views respond in under two seconds under normal pilot load.",
    "Retries do not create duplicate marks or student records.",
    "Every rejected import row shows an actionable reason.",
    "The school can export its records and Nexatech can demonstrate a successful backup restore.",
]:
    bullet(doc, item)

doc.add_heading("System Dependencies and Libraries", level=1)
body(doc, "The list below covers direct project dependencies and the main production runtime components. Package managers also install transitive dependencies required by these libraries.")

doc.add_heading("Frontend Application", level=2)
add_table(doc,
    ["Component", "Version", "Purpose"],
    [
        ("React", "19.2.6", "User interface and role-specific application screens"),
        ("React DOM", "19.2.6", "Browser rendering for React"),
        ("Vite", "8.0.14", "Development server and production frontend build"),
        ("Tailwind CSS", "4.3.0", "Layout and visual styling"),
        ("Tailwind Vite plugin", "4.3.0", "Tailwind integration with the build"),
        ("Axios", "1.16.1", "Authenticated browser requests to the Laravel API"),
        ("Playwright", "1.60.0", "Browser testing and repeatable demonstrations"),
        ("ESLint toolchain", "10.x", "JavaScript quality checks, React hooks rules, and development globals"),
        ("React type packages", "19.2.x", "Development type information"),
    ],
    [1.9, 1.1, 4.15]
)

doc.add_heading("Backend Application", level=2)
add_table(doc,
    ["Component", "Version", "Purpose"],
    [
        ("PHP", "8.3 minimum; 8.4 production image", "Server runtime"),
        ("Laravel Framework", "13.8", "API, authentication, validation, queues, scheduling, and business logic"),
        ("Laravel Sanctum", "4.0", "Authenticated sessions and access tokens"),
        ("Laravel Tinker", "3.0", "Controlled administrative troubleshooting"),
        ("Laravel DOMPDF", "3.1", "PDF report-card generation"),
        ("Paynow PHP SDK", "1.0", "Paynow payment integration"),
        ("PHP extensions", "bcmath, gd, intl, pdo_mysql, zip, redis", "Calculations, images, localization, database, archives, queues, and cache"),
        ("Test and development libraries", "Faker, Pail, Pao, Pint, Mockery, Collision, PHPUnit", "Fixtures, logs, quality checks, mocks, diagnostics, and automated tests"),
    ],
    [2.0, 1.65, 3.5]
)

doc.add_heading("System and Service Dependencies", level=2)
add_table(doc,
    ["Dependency", "Use"],
    [
        ("Docker", "Builds one repeatable production image"),
        ("Node.js 22 and npm", "Build the React frontend"),
        ("Composer 2", "Installs and optimizes PHP packages"),
        ("Apache", "Serves the built frontend and Laravel application"),
        ("Managed MySQL", "Production source of truth for school records"),
        ("Redis", "Queues, cache, and session support"),
        ("SMTP service", "Password recovery and transactional email"),
        ("Africa's Talking", "Optional SMS delivery"),
        ("Paynow, Payfast, or Stitch", "Optional payment processing through configured gateways"),
    ],
    [2.15, 5.0]
)

doc.add_heading("Hosting and Production Architecture", level=1)
body(doc, "The current deployment design uses Railway and Docker. The React application and Laravel API share one HTTPS origin because the secure authentication design uses HTTP-only cookies.")
add_table(doc,
    ["Service", "Responsibility"],
    [
        ("Web service", "Serves the React frontend and Laravel API, runs migrations, and exposes the health check"),
        ("Worker service", "Processes queued email, SMS, payment confirmation, and other background jobs"),
        ("Scheduler service", "Runs scheduled work such as retention pruning"),
        ("Managed MySQL", "Stores tenant-scoped operational data"),
        ("Managed Redis", "Supports queues, cache, and production sessions"),
        ("Custom domain and HTTPS", "Provides the school-facing address and encrypted transport"),
        ("Monitoring and backups", "Tracks health, failures, database pressure, and recovery readiness"),
    ],
    [2.0, 5.15]
)

doc.add_heading("Domain API and Secret Management", level=1)
for item in [
    "Use one final product domain with HTTPS and domain-based sender addresses.",
    "Store application, database, email, SMS, payment, and partner credentials in the hosting secret manager.",
    "Never place API keys in the React frontend, screenshots, source control, proposals, or training material.",
    "Issue a unique key for each customer integration with narrow permissions, rate limits, rotation, logs, and revocation.",
    "Charge for managed API access and usage, not for revealing a provider secret.",
]:
    bullet(doc, item)

doc.add_heading("Readiness Conditions Before Real School Data", level=1)
for item in [
    "Complete tenant-isolation and role-permission checks for the exact pilot roles.",
    "Confirm automated encrypted backups and complete a documented restore test.",
    "Activate uptime, application-error, queue-failure, and database alerts.",
    "Configure and test SMTP before relying on password recovery or demo access email.",
    "Configure SMS and payment credentials only for services included in the pilot.",
    "Approve the privacy notice, terms, retention policy, data-processing responsibilities, and incident contact process.",
    "Use synthetic data for sales demonstrations. Import real school data only after written approval and secure transfer arrangements.",
]:
    bullet(doc, item)

doc.add_heading("Next Decision", level=1)
body(doc, "The school and Nexatech should agree the pilot workflow, nominate one school lead, approve the data set, confirm the price, and schedule a 30-minute demonstration. The pilot should begin only after both sides sign the scope and data-handling responsibilities.")

doc.core_properties.title = "Nexatech School Pilot Proposal and Budget"
doc.core_properties.subject = "School onboarding proposal, demonstration plan, pricing, dependencies, and hosting"
doc.core_properties.author = "NEXATECH"
doc.core_properties.keywords = "Nexatech, school onboarding, proposal, budget, demo, hosting, dependencies"

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
