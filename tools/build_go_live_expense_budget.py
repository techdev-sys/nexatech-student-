from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "deliverables" / "NEXATECH Go Live Expense Budget.docx"

NAVY = "16324F"
PALE_BLUE = "EDF4F8"
LIGHT_GRAY = "F5F6F7"
BORDER = "D9D9D9"
TEXT = "1A1A1A"
MUTED = "5F6B76"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), BORDER)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Aptos", size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.08):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def remove_paragraph_borders(paragraph_or_style):
    element = paragraph_or_style._element
    p_pr = element.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def style_cell_text(cell, bold=False, color=TEXT, align=WD_ALIGN_PARAGRAPH.LEFT, size=9.5):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)
    for paragraph in cell.paragraphs:
        paragraph.alignment = align
        set_paragraph_spacing(paragraph, after=0, line=1.05)
        for run in paragraph.runs:
            set_run_font(run, size=size, bold=bold, color=color)


def add_expense_table(doc, rows, columns):
    table = doc.add_table(rows=1, cols=len(columns))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    widths = [Inches(c[1]) for c in columns]

    header = table.rows[0]
    set_repeat_table_header(header)
    for index, (label, _) in enumerate(columns):
        cell = header.cells[index]
        cell.width = widths[index]
        cell.text = label
        set_cell_shading(cell, NAVY)
        align = WD_ALIGN_PARAGRAPH.RIGHT if index == len(columns) - 1 else WD_ALIGN_PARAGRAPH.LEFT
        style_cell_text(cell, bold=True, color="FFFFFF", align=align, size=9)

    for row_index, values in enumerate(rows):
        row = table.add_row()
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.width = widths[index]
            cell.text = str(value)
            if row_index % 2:
                set_cell_shading(cell, LIGHT_GRAY)
            align = WD_ALIGN_PARAGRAPH.RIGHT if index == len(columns) - 1 else WD_ALIGN_PARAGRAPH.LEFT
            style_cell_text(cell, align=align)

    return table


def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_run_font(first, size=10.5, bold=True, color=TEXT)
        second = p.add_run(text[len(bold_prefix):])
        set_run_font(second, size=10.5, color=TEXT)
    else:
        run = p.add_run(text)
        set_run_font(run, size=10.5, color=TEXT)
    set_paragraph_spacing(p, after=3, line=1.1)
    return p


def add_source(doc, number, name, url, detail):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    set_paragraph_spacing(p, after=5, line=1.05)
    lead = p.add_run(f"{number}. {name}. ")
    set_run_font(lead, size=8.5, bold=True, color=TEXT)
    body = p.add_run(f"{detail} {url}")
    set_run_font(body, size=8.5, color=MUTED)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(TEXT)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.1

title_style = styles["Title"]
title_style.font.name = "Aptos Display"
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
title_style.font.size = Pt(25)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_style.paragraph_format.space_after = Pt(8)
remove_paragraph_borders(title_style)

for style_name, size, before, after in (("Heading 1", 16, 13, 6), ("Heading 2", 12, 10, 4)):
    style = styles[style_name]
    style.font.name = "Aptos Display"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hr = hp.add_run("NEXATECH STUDENT MONITOR")
set_run_font(hr, size=8, bold=True, color=MUTED)

footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = fp.add_run("Go Live Expense Budget  |  21 October 2026  |  USD")
set_run_font(fr, size=8, color=MUTED)

title_paragraph = doc.add_paragraph("NEXATECH Student Monitor Go Live Expense Budget", style="Title")
remove_paragraph_borders(title_paragraph)

subtitle = doc.add_paragraph()
set_paragraph_spacing(subtitle, after=12)
sr = subtitle.add_run("Approved launch package for one school")
set_run_font(sr, size=12, bold=True, color=NAVY)

intro = doc.add_paragraph()
set_paragraph_spacing(intro, after=8, line=1.12)
run = intro.add_run(
    "This budget assigns exact amounts for launching and operating NEXATECH Student Monitor for one school. "
    "The approved funding is US$425.40 for the first month and US$60.40 for each normal operating month after launch."
)
set_run_font(run, size=10.8, color=TEXT)

scope = doc.add_paragraph()
set_paragraph_spacing(scope, after=10)
scope_lead = scope.add_run("Fixed scope: ")
set_run_font(scope_lead, size=10.5, bold=True, color=TEXT)
scope_body = scope.add_run(
    "one production system, no more than 1,000 total data subjects, one business mailbox, up to 3,000 transactional emails per month, and 500 SMS messages per month."
)
set_run_font(scope_body, size=10.5, color=TEXT)

doc.add_heading("Approved Monthly Operating Expenses", level=1)
monthly_rows = [
    ("Production hosting", "Railway application, worker, MySQL and Redis", "$40.00"),
    ("Business mailbox", "One Google Workspace Business Starter mailbox", "$8.40"),
    ("Transactional email", "Resend Free plan, maximum 3,000 emails per month", "$0.00"),
    ("SMS messaging", "Prepaid allowance for 500 local SMS messages", "$10.00"),
    ("Database recovery", "Backup and recovery storage allocation", "$2.00"),
    ("Uptime monitoring", "Better Stack Free plan", "$0.00"),
    ("CAPTCHA protection", "Cloudflare Turnstile Free plan", "$0.00"),
    ("SSL and DNS", "Cloudflare and Railway certificates and DNS", "$0.00"),
    ("Monthly total", "Approved recurring operating budget", "$60.40"),
]
table = add_expense_table(
    doc,
    monthly_rows,
    (("Expense", 1.65), ("Included service", 4.25), ("Monthly USD", 1.15)),
)
for index, cell in enumerate(table.rows[-1].cells):
    set_cell_shading(cell, PALE_BLUE)
    style_cell_text(
        cell,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.RIGHT if index == len(table.rows[-1].cells) - 1 else WD_ALIGN_PARAGRAPH.LEFT,
    )

doc.add_heading("Launch and Annual Expenses", level=1)
launch_rows = [
    ("Domain registration", "Annual", "$15.00"),
    ("POTRAZ Tier 1 data controller licence", "Annual", "$50.00"),
    ("Zimbabwe privacy and contract review", "Launch only", "$300.00"),
    ("Launch expense total", "Due before public onboarding", "$365.00"),
]
table2 = add_expense_table(
    doc,
    launch_rows,
    (("Expense", 4.15), ("Frequency", 1.75), ("Amount USD", 1.15)),
)
for index, cell in enumerate(table2.rows[-1].cells):
    set_cell_shading(cell, PALE_BLUE)
    style_cell_text(
        cell,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.RIGHT if index == len(table2.rows[-1].cells) - 1 else WD_ALIGN_PARAGRAPH.LEFT,
    )

doc.add_page_break()

doc.add_heading("Budget Totals", level=1)
summary_rows = [
    ("First month cash requirement", "$365.00 launch expenses plus $60.40 operations", "$425.40"),
    ("Months 2 through 12", "$60.40 each month for 11 months", "$664.40"),
    ("First year total", "Complete launch year", "$1,089.80"),
    ("Year 2 baseline", "$724.80 operations plus $65 annual renewals", "$789.80"),
]
table3 = add_expense_table(
    doc,
    summary_rows,
    (("Budget period", 2.05), ("Calculation", 3.85), ("Total USD", 1.15)),
)
for target_row in (table3.rows[1], table3.rows[3], table3.rows[4]):
    for index, cell in enumerate(target_row.cells):
        if target_row in (table3.rows[3], table3.rows[4]):
            set_cell_shading(cell, PALE_BLUE)
            style_cell_text(
                cell,
                bold=True,
                align=WD_ALIGN_PARAGRAPH.RIGHT if index == len(target_row.cells) - 1 else WD_ALIGN_PARAGRAPH.LEFT,
            )

doc.add_heading("Financial Rules", level=1)
add_bullet(doc, "Hosting allocation: Railway spending is limited to the approved $40.00 monthly infrastructure budget.", "Hosting allocation:")
add_bullet(doc, "SMS limit: The system may send no more than 500 prepaid SMS messages in a month without a new approval.", "SMS limit:")
add_bullet(doc, "Payment fees: Paynow transaction fees are added to the payer's transaction and are not charged to this operating budget.", "Payment fees:")
add_bullet(doc, "Email limit: Resend remains on the free plan and must stay within 3,000 emails per month and 100 emails per day.", "Email limit:")
add_bullet(doc, "Scope control: A second school or more than 1,000 data subjects requires a new operating budget and licensing review.", "Scope control:")
add_bullet(doc, "No hidden contingency: Spending above any line in this document requires written approval before purchase.", "No hidden contingency:")

doc.add_heading("What This Budget Delivers", level=1)
p = doc.add_paragraph()
set_paragraph_spacing(p, after=8, line=1.12)
r = p.add_run(
    "The approved package provides a custom domain, HTTPS, production hosting, a MySQL database, Redis queues, "
    "scheduled processing, one professional email inbox, transactional email, 500 monthly SMS messages, backups, "
    "monitoring, CAPTCHA protection, and launch compliance review. It does not include additional schools, paid advertising, "
    "custom mobile applications, staff devices, internet connectivity, or payment transaction charges."
)
set_run_font(r, size=10.5, color=TEXT)

doc.add_heading("Pricing Sources", level=1)
add_source(doc, 1, "Railway", "https://docs.railway.com/pricing/plans", "Pro plan and usage-based infrastructure pricing.")
add_source(doc, 2, "Google Workspace", "https://knowledge.workspace.google.com/admin/billing/compare-flexible-and-annual-fixed-term-payment-plans", "Business Starter flexible plan at $8.40 per user per month.")
add_source(doc, 3, "Resend", "https://resend.com/pricing", "Free transactional email plan with 3,000 emails per month and a 100-email daily limit.")
add_source(doc, 4, "SMSPoP Zimbabwe", "https://www.smspop.co.zw/pricing", "Prepaid package of 500 SMS credits for $10.")
add_source(doc, 5, "Cloudflare", "https://developers.cloudflare.com/registrar/", "At-cost domain registration, DNS and SSL services.")
add_source(doc, 6, "Better Stack", "https://betterstack.com/uptime", "Free uptime monitors and heartbeat monitoring.")
add_source(doc, 7, "POTRAZ", "https://www.potraz.gov.zw/wp-content/uploads/2025/02/sI-155-of-2024-Cyber-and-Data-Protection-Normal_240913_1250178.pdf", "Tier 1 data controller licence fee of $50.")
add_source(doc, 8, "Paynow", "https://www.paynow.co.zw/Home/Fees", "Transaction fees can be added to the payer rather than absorbed by the platform.")

doc.core_properties.title = "NEXATECH Student Monitor Go Live Expense Budget"
doc.core_properties.subject = "Fixed launch and operating budget"
doc.core_properties.author = "NEXATECH"
doc.core_properties.keywords = "NEXATECH, expenses, hosting, deployment, budget"

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
